import io
import xml.etree.ElementTree as ET
import zipfile
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import MagicMock, patch

import docx
import pytest
from fastapi.testclient import TestClient

from app.config import settings
from app.extractors.docx import DocxExtractor
from app.extractors.pages import PagesExtractor
from app.main import app


def create_mock_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Test Heading 1", level=1)
    doc.add_paragraph("This is paragraph 1 inside the document.")
    doc.add_paragraph("This is paragraph 2.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header A"
    table.cell(0, 1).text = "Header B"
    table.cell(1, 0).text = "Cell A1"
    table.cell(1, 1).text = "Cell B1"

    doc.add_paragraph("Another paragraph after table.")

    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def create_docx_with_content_controls() -> bytes:
    doc = docx.Document()
    doc.add_paragraph("Title inside content control")
    doc.add_paragraph("Body outside content control")
    doc.add_paragraph("Reference inside content control")

    stream = io.BytesIO()
    doc.save(stream)

    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ET.register_namespace("w", ns)

    source = zipfile.ZipFile(io.BytesIO(stream.getvalue()))
    document_xml = source.read("word/document.xml")
    root = ET.fromstring(document_xml)
    body = root.find(f"{{{ns}}}body")
    assert body is not None

    paragraph_indexes = [i for i, child in enumerate(list(body)) if child.tag == f"{{{ns}}}p"]

    for child_index in (paragraph_indexes[-1], paragraph_indexes[0]):
        child = list(body)[child_index]
        sdt = ET.Element(f"{{{ns}}}sdt")
        sdt_content = ET.SubElement(sdt, f"{{{ns}}}sdtContent")
        body.remove(child)
        sdt_content.append(child)
        body.insert(child_index, sdt)

    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as target:
        for item in source.infolist():
            data = (
                ET.tostring(root, encoding="utf-8", xml_declaration=True)
                if item.filename == "word/document.xml"
                else source.read(item.filename)
            )
            target.writestr(item, data)

    source.close()
    return output.getvalue()


def create_mock_zip(files: dict[str, bytes]) -> bytes:
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as zf:
        for filename, data in files.items():
            zf.writestr(filename, data)
    return stream.getvalue()


@pytest.mark.asyncio
async def test_docx_extractor_correctness():
    executor = ThreadPoolExecutor(max_workers=1)
    extractor = DocxExtractor(executor)

    content = create_mock_docx()
    text = await extractor.extract(content, filename="test.docx")

    assert "# Test Heading 1" in text
    assert "This is paragraph 1 inside the document." in text
    assert "| Header A | Header B |" in text
    assert "| Cell A1 | Cell B1 |" in text
    assert "Another paragraph after table." in text
    executor.shutdown()


@pytest.mark.asyncio
async def test_docx_extractor_reads_content_controls():
    executor = ThreadPoolExecutor(max_workers=1)
    extractor = DocxExtractor(executor)

    content = create_docx_with_content_controls()
    text = await extractor.extract(content, filename="content_controls.docx")

    assert "Title inside content control" in text
    assert "Body outside content control" in text
    assert "Reference inside content control" in text
    assert text.index("Title inside content control") < text.index("Body outside content control")
    assert text.index("Body outside content control") < text.index(
        "Reference inside content control"
    )
    executor.shutdown()


@pytest.mark.asyncio
async def test_pages_extractor_strategy_1_pdf():
    # Test Strategy 1 (post-2013 Pages) - PDF preview inside zip
    executor = ThreadPoolExecutor(max_workers=1)
    extractor = PagesExtractor(executor)

    # Mock PDF extractor and its _sync_extract method
    mock_pdf_extractor = MagicMock()
    mock_pdf_extractor._sync_extract.return_value = "Extracted PDF Text from Pages"
    extractor._pdf_extractor = mock_pdf_extractor

    # Create pages mock content which is zip containing QuickLook/Preview.pdf
    mock_pdf_bytes = b"%PDF-1.4 mock content"
    content = create_mock_zip({"QuickLook/Preview.pdf": mock_pdf_bytes})

    text = await extractor.extract(content, filename="test.pages")

    # Assert that signature mismatch does not occur and correct return is got
    assert text == "Extracted PDF Text from Pages"
    mock_pdf_extractor._sync_extract.assert_called_once_with(mock_pdf_bytes)
    executor.shutdown()


@pytest.mark.asyncio
async def test_pages_extractor_strategy_2_xml():
    # Test Strategy 2 (pre-2013 Pages) - index.xml
    executor = ThreadPoolExecutor(max_workers=1)
    extractor = PagesExtractor(executor)

    xml_content = b"""<?xml version="1.0" encoding="utf-8"?>
    <index>
      <text>Some text in index.xml</text>
      <tail>Tail text</tail>
    </index>
    """
    content = create_mock_zip({"index.xml": xml_content})

    text = await extractor.extract(content, filename="test_old.pages")

    # Assert that text was successfully parsed out of XML
    assert "Some text in index.xml" in text
    assert "Tail text" in text
    executor.shutdown()


def test_api_health():
    with TestClient(app) as client:
        response = client.get("/health")
        assert response.status_code == 200
        assert response.json() == {"status": "ok"}


def test_api_unauthorized_extract():
    # If API_KEY is set, make sure request without header is unauthorized
    if not settings.API_KEY:
        return

    with TestClient(app) as client:
        response = client.post("/extract/urls", json={"urls": []})
        assert response.status_code == 401


@pytest.mark.asyncio
async def test_image_extractor():
    from app.extractors.image import ImageExtractor

    executor = ThreadPoolExecutor(max_workers=1)
    extractor = ImageExtractor(executor)

    with patch("PIL.Image.open") as mock_open, patch("pytesseract.image_to_string") as mock_ocr:
        mock_ocr.return_value = "Mocked OCR Text Results"

        text = await extractor.extract(b"mock image bytes", filename="test.png")

        assert text == "Mocked OCR Text Results"
        mock_open.assert_called_once()
        mock_ocr.assert_called_once()

    executor.shutdown()


@pytest.mark.asyncio
async def test_plaintext_extractor_instantiation_and_extraction():
    from app.extractors.plaintext import PlainTextExtractor

    executor = ThreadPoolExecutor(max_workers=1)
    extractor = PlainTextExtractor(executor)

    raw = b"# Sample Heading\nThis is plain text."
    result = await extractor.extract(raw, filename="test.md")
    assert "# Sample Heading" in result
    assert "This is plain text." in result
    executor.shutdown()


def test_cache_rejects_blank_content():
    from app.routers.url_extract import URLExtractResult
    from app.utils.cache import URLCache

    cache = URLCache(ttl_seconds=60)

    # Empty string should not be cached
    cache.set(
        "https://example.com/empty", URLExtractResult(url="https://example.com/empty", content="")
    )
    assert cache.get("https://example.com/empty") is None

    # Whitespace-only string should not be cached
    cache.set(
        "https://example.com/ws", URLExtractResult(url="https://example.com/ws", content="\n  \n\t")
    )
    assert cache.get("https://example.com/ws") is None

    # Meaningful text should be cached
    cache.set(
        "https://example.com/valid",
        URLExtractResult(url="https://example.com/valid", content="Meaningful extracted text"),
    )
    cached = cache.get("https://example.com/valid")
    assert cached is not None
    assert cached.content == "Meaningful extracted text"


def test_url_extract_request_coercion():
    from app.routers.url_extract import URLExtractRequest

    # Single string URL should be coerced into list
    req = URLExtractRequest.model_validate({"urls": "https://example.com/document.pdf"})
    assert isinstance(req.urls, list)
    assert len(req.urls) == 1
    assert str(req.urls[0]) == "https://example.com/document.pdf"

    # List of URLs should remain list
    req2 = URLExtractRequest.model_validate(
        {"urls": ["https://example.com/1.pdf", "https://example.com/2.pdf"]}
    )
    assert len(req2.urls) == 2


def test_cache_endpoints():
    with TestClient(app) as client:
        # Initial stats
        res = client.get("/cache")
        assert res.status_code == 200
        data = res.json()
        assert "cached_entries" in data
        assert "ttl_seconds" in data

        # Flush cache
        del_res = client.delete("/cache")
        assert del_res.status_code == 200
        assert "flushed" in del_res.json()


@pytest.mark.asyncio
async def test_pdf_ocr_fallback_triggers_when_empty():
    from app.extractors.pdf import PDFExtractor

    executor = ThreadPoolExecutor(max_workers=1)
    extractor = PDFExtractor(executor)

    call_history = []

    def mock_to_markdown(doc, use_ocr=False, **kwargs):
        call_history.append({"use_ocr": use_ocr, **kwargs})
        if not use_ocr:
            return "\n\n\n"  # Pass 1 returns blank text
        return "# OCR Extracted Title\nSome OCR content"  # Pass 2 returns OCR content

    with (
        patch("pymupdf4llm.to_markdown", side_effect=mock_to_markdown),
        patch("pymupdf.open") as mock_open,
    ):
        mock_doc = MagicMock()
        mock_doc.is_encrypted = False
        mock_doc.__len__.return_value = 1
        mock_doc.__enter__.return_value = mock_doc
        mock_doc.__exit__.return_value = False
        mock_open.return_value = mock_doc

        result = await extractor.extract(b"%PDF mock", filename="test.pdf")

    # Assert both passes were called
    assert len(call_history) == 2
    assert call_history[0]["use_ocr"] is False
    assert call_history[1]["use_ocr"] is True
    assert "OCR Extracted Title" in result

    executor.shutdown()


@pytest.mark.asyncio
async def test_pdf_encrypted_raises_value_error():
    from app.extractors.pdf import PDFExtractor

    executor = ThreadPoolExecutor(max_workers=1)
    extractor = PDFExtractor(executor)

    with patch("pymupdf.open") as mock_open:
        mock_doc = MagicMock()
        mock_doc.is_encrypted = True
        mock_doc.__enter__.return_value = mock_doc
        mock_doc.__exit__.return_value = False
        mock_open.return_value = mock_doc

        with pytest.raises(ValueError, match="encrypted"):
            await extractor.extract(b"%PDF mock", filename="test.pdf")

    executor.shutdown()


def test_binary_extraction_endpoint():
    import json

    with TestClient(app) as client:
        metadata = [{"binaryKey": "file1", "fileName": "notes.txt", "docId": 42}]
        files = {"file1": ("notes.txt", b"Extracted from multipart form upload", "text/plain")}
        res = client.post("/extract/binary", data={"data": json.dumps(metadata)}, files=files)
        assert res.status_code == 200
        data = res.json()
        assert len(data) == 1
        assert data[0]["text"] == "Extracted from multipart form upload"
        assert data[0]["docId"] == 42


def test_binary_extraction_with_surrogates():
    import json

    with TestClient(app) as client:
        with patch(
            "app.extractors.pdf.PDFExtractor.extract",
            return_value="Title with surrogate \ud835\ud835 text",
        ):
            metadata = [{"binaryKey": "file1", "fileName": "31.pdf"}]
            files = {"file1": ("31.pdf", b"%PDF-mock", "application/pdf")}
            res = client.post("/extract/binary", data={"data": json.dumps(metadata)}, files=files)
            assert res.status_code == 200
            data = res.json()
            assert data[0]["text"] == "Title with surrogate  text"


def test_mime_type_detection_without_extension():
    from app.utils.format_router import detect_format

    # No extension in URL, but MIME type is provided
    assert detect_format(b"hello world", filename="https://example.com/item/123", content_type="text/html") == ".html"
    assert detect_format(b"col1,col2\n1,2", filename="https://example.com/export?id=99", content_type="text/csv") == ".csv"
    assert detect_format(b"%PDF-1.4", filename="https://example.com/api/get", content_type="application/pdf") == ".pdf"


def test_sanitize_surrogates_and_nul_bytes():
    from app.utils.sanitizer import sanitize_surrogates

    assert sanitize_surrogates("valid text\x00with nul") == "valid textwith nul"
    assert sanitize_surrogates("surrogate \ud835\ud835 test") == "surrogate  test"


@pytest.mark.asyncio
async def test_pdf_fallback_to_raw_text_on_layout_error():
    from app.extractors.pdf import PDFExtractor

    executor = ThreadPoolExecutor(max_workers=1)
    extractor = PDFExtractor(executor)

    with patch("pymupdf.open") as mock_open:
        mock_page = MagicMock()
        mock_page.get_text.return_value = "Fallback raw text from PyMuPDF"

        mock_doc = MagicMock()
        mock_doc.is_encrypted = False
        mock_doc.__len__.return_value = 1
        mock_doc.__iter__.return_value = [mock_page]
        mock_doc.__enter__.return_value = mock_doc
        mock_doc.__exit__.return_value = False
        mock_open.return_value = mock_doc

        with patch("pymupdf4llm.to_markdown", side_effect=RuntimeError("Layout parsing model crashed")):
            text = await extractor.extract(b"%PDF mock", filename="test.pdf")
            assert "Fallback raw text from PyMuPDF" in text

    executor.shutdown()


@pytest.mark.asyncio
async def test_url_extraction_inflight_coalescing():
    import asyncio

    from app.routers.url_extract import URLExtractResult, _inflight_tasks, _process_url

    call_count = 0

    async def mock_execute(*args, **kwargs):
        nonlocal call_count
        call_count += 1
        await asyncio.sleep(0.05)
        return URLExtractResult(url="https://example.com/shared.pdf", content="Extracted once", error=None)

    with patch("app.routers.url_extract.validate_safe_url"):
        with patch("app.routers.url_extract.url_cache.get", return_value=None):
            with patch("app.routers.url_extract._execute_url_extraction", side_effect=mock_execute):
                sem = asyncio.Semaphore(5)
                r1, r2 = await asyncio.gather(
                    _process_url("https://example.com/shared.pdf", None, None, sem, sem),
                    _process_url("https://example.com/shared.pdf", None, None, sem, sem),
                )
                assert r1.content == "Extracted once"
                assert r2.content == "Extracted once"
                assert call_count == 1
                assert len(_inflight_tasks) == 0


